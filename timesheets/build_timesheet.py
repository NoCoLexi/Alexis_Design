"""Build a monthly timesheet (deliverable rows). Usage: python build_timesheet.py <upstart|sikich> <YYYY> <M>"""
import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from profiles import PROFILES

def build(key, year, month, out=None):
    P = PROFILES[key]
    fill = lambda c: PatternFill("solid", fgColor=c)
    F = lambda n=None, s=10, b=False, c=None, i=False: Font(name=n or P["sans"], size=s, bold=b, color=c or P["ink"], italic=i)
    hair = Side(style="thin", color=P["line"]); box = Border(left=hair, right=hair, top=hair, bottom=hair)
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lft = Alignment(horizontal="left", vertical="center", wrap_text=True)

    days, d = [], dt.date(year, month, 1)
    while d.month == month:
        if d.weekday() < 5: days.append(d)
        d += dt.timedelta(1)
    n = len(days); D0 = 4; D1 = D0 + n - 1
    TOT, DAYS, BUD, PRI, PCT, STAT, NOTE = (D1 + i for i in range(1, 8))
    HR, R0, NROWS = 5, 6, 8; R1 = R0 + NROWS - 1; TR = R1 + 2

    wb = Workbook(); ws = wb.active; ws.title = "Timesheet"; ws.sheet_view.showGridLines = False
    for r in range(1, TR + 12):
        for c in range(1, NOTE + 2): ws.cell(r, c).fill = fill(P["page"])
    widths = {1: 2, 2: 34, 3: 9, TOT: 9, DAYS: 8, BUD: 11, PRI: 11, PCT: 11, STAT: 14, NOTE: 40}
    for c in range(1, NOTE + 1): ws.column_dimensions[L(c)].width = widths.get(c, 5.5)

    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=NOTE)
    ws["B1"] = f"{P['title']}: {dt.date(year, month, 1):%B %Y}"
    ws["B1"].font = F(P["serif"], 18 if key == "upstart" else 14, key == "sikich", P["title_text"])
    ws["B1"].alignment = Alignment(vertical="center"); ws.row_dimensions[1].height = 32
    if P["title_fill"]:
        for c in range(2, NOTE + 1): ws.cell(1, c).fill = fill(P["title_fill"])
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=NOTE)
    ws["B2"] = f"{P['tag']}  |  {n} working days"
    ws["B2"].font = F(P["mono"], 9, False, P["muted"]); ws.row_dimensions[2].height = 18
    if key == "sikich":
        for c in range(2, NOTE + 1): ws.cell(2, c).fill = fill(P["band"])
    ws.row_dimensions[3].height = 3
    if key == "upstart": ws.cell(3, 2).fill = fill(P["accent"])

    ws["B4"] = "Hours per day (1.0 FTE basis)"; ws["B4"].font = F(P["mono"], 8, False, P["muted"]); ws["B4"].alignment = lft
    ws["C4"] = 8; ws["C4"].fill = fill(P["input"]); ws["C4"].font = F(None, 10, True); ws["C4"].alignment = ctr; ws["C4"].border = box
    HPD = "$C$4"

    weeks = {}
    for i, dd in enumerate(days): weeks.setdefault(dd.isocalendar()[1], []).append(D0 + i)
    for k, cols in enumerate(weeks.values(), 1):
        a, b = cols[0], cols[-1]
        if a != b: ws.merge_cells(start_row=4, start_column=a, end_row=4, end_column=b)
        ws.cell(4, a, f"Week {k}").font = F(P["mono"], 8, True); ws.cell(4, a).alignment = ctr
        for c in range(a, b + 1): ws.cell(4, c).fill = fill(P["soft"]); ws.cell(4, c).border = box

    heads = {2: "Deliverable / artifact", 3: "Phase", TOT: "Total\nhours", DAYS: "Days", BUD: "SOW budget\n(hrs)",
             PRI: "Billed prior\n(hrs)", PCT: "% of SOW\nused", STAT: "Status", NOTE: "Notes"}
    for c, t in heads.items(): ws.cell(HR, c, t)
    for i, dd in enumerate(days): ws.cell(HR, D0 + i, f"{dd:%a}\n{dd.day}")
    for c in range(2, NOTE + 1):
        x = ws.cell(HR, c); x.font = F(None, 8 if D0 <= c <= D1 else 10, True, P["head_text"]); x.fill = fill(P["head_fill"])
        x.alignment = lft if c in (2, NOTE) else ctr; x.border = box
    ws.row_dimensions[HR].height = 32

    for k in range(NROWS):
        r = R0 + k
        for c in range(2, NOTE + 1):
            x = ws.cell(r, c); x.border = box; x.alignment = ctr; x.font = F()
            x.fill = fill(P["input"])
        ws.cell(r, 2).alignment = lft; ws.cell(r, NOTE).alignment = lft; ws.cell(r, NOTE).font = F(None, 9, False, P["muted"])
        for c in range(D0, D1 + 1): ws.cell(r, c).number_format = "0.0;;"
        ws.cell(r, BUD).number_format = ws.cell(r, PRI).number_format = "0.0;;"
        for c, fm, nf in ((TOT, f"=SUM({L(D0)}{r}:{L(D1)}{r})", "0.0"), (DAYS, f"={L(TOT)}{r}/{HPD}", "0.00"),
                          (PCT, f'=IFERROR(({L(PRI)}{r}+{L(TOT)}{r})/{L(BUD)}{r},"")', "0%")):
            x = ws.cell(r, c, fm); x.fill = fill(P["band"]); x.number_format = nf; x.font = F(None, 10, c == TOT)
        ws.row_dimensions[r].height = 22
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", allow_blank=True,
                        errorTitle="Check hours", error="Enter a number from 0 to 24.", showErrorMessage=True)
    dv.add(f"{L(D0)}{R0}:{L(D1)}{R1}")
    ds = DataValidation(type="list", formula1='"Not started,In progress,In review,Complete"', allow_blank=True)
    ds.add(f"{L(STAT)}{R0}:{L(STAT)}{R1}"); ws.add_data_validation(dv); ws.add_data_validation(ds)

    ws.merge_cells(start_row=TR, start_column=2, end_row=TR, end_column=3); ws.cell(TR, 2, "Monthly totals")
    for c in range(2, NOTE + 1):
        x = ws.cell(TR, c); x.fill = fill(P["soft"]); x.border = box; x.font = F(None, 10, True); x.alignment = ctr
    ws.cell(TR, 2).alignment = lft
    for c, nf in [(c, "0.0") for c in range(D0, TOT + 1)] + [(DAYS, "0.00"), (BUD, "0.0"), (PRI, "0.0")]:
        ws.cell(TR, c, f"=SUM({L(c)}{R0}:{L(c)}{R1})").number_format = nf
    ws.cell(TR, PCT, f'=IFERROR(({L(PRI)}{TR}+{L(TOT)}{TR})/{L(BUD)}{TR},"")').number_format = "0%"
    ws.cell(TR, TOT).font = F(None, 11, True, P["emph"]); ws.cell(TR, PCT).font = F(None, 11, True, P["emph"])
    ws.row_dimensions[TR].height = 26

    lr = TR + 2
    ws.merge_cells(start_row=lr, start_column=2, end_row=lr, end_column=NOTE); ws.cell(lr, 2, "How to use").font = F(P["serif"], 12, key == "sikich")
    leg = [f"Enter hours by day against each deliverable in the {'white' if key == 'upstart' else 'yellow'} cells. Everything else calculates.",
           "Phase is the SOW phase (for example Phase I). SOW budget is the hours the SOW allows for that deliverable.",
           "Billed prior is hours already billed in earlier months, so % of SOW used reflects position to date.",
           "Example, not counted in totals: Change management plan | Phase I | Thu 1 = 6.5, Fri 2 = 4 | SOW budget 80 | Billed prior 20 | Status In progress."]
    for i, t in enumerate(leg, 1):
        ws.merge_cells(start_row=lr + i, start_column=2, end_row=lr + i, end_column=NOTE)
        ws.cell(lr + i, 2, t).font = F(None, 9, False, P["muted"], i == 4)
    ws.freeze_panes = ws.cell(R0, D0)
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = out or f"{P['file_prefix']}_Timesheet_{year}-{month:02d}.xlsx"
    wb.save(out); return out

if __name__ == "__main__":
    k, y, m = sys.argv[1], int(sys.argv[2]), int(sys.argv[3]); print(build(k, y, m))
