"""One-page Word executive summary from a filled timesheet.
Usage: python build_summary.py <upstart|sikich> <timesheet.xlsx> [narrative.json]
narrative.json (optional): {"phase": "Phase I", "risks": ["..."], "next": ["..."]}
Inputs are read from the timesheet and totals are recomputed here, so no cached Excel values are needed."""
import sys, json, re
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from profiles import PROFILES

def read(path):
    ws = load_workbook(path).active
    hdr = {str(ws.cell(5, c).value).replace("\n", " "): c for c in range(2, ws.max_column + 1) if ws.cell(5, c).value}
    days = [c for c in range(4, ws.max_column + 1) if re.match(r"^[A-Z][a-z]{2}\n\d+$", str(ws.cell(5, c).value or ""))]
    hpd = ws["C4"].value or 8; rows = []
    for r in range(6, 14):
        name = ws.cell(r, 2).value
        if not name: continue
        hrs = sum(ws.cell(r, c).value or 0 for c in days)
        rows.append(dict(name=name, phase=ws.cell(r, 3).value or "", hrs=hrs,
                         bud=ws.cell(r, hdr["SOW budget (hrs)"]).value or 0,
                         prior=ws.cell(r, hdr["Billed prior (hrs)"]).value or 0,
                         status=ws.cell(r, hdr["Status"]).value or "Not started"))
    return ws["B1"].value, len(days), hpd, rows

def shade(cell, hexc):
    p = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexc); p.append(s)

def borders(cell, hexc, sides=("bottom",)):
    p = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}"); e.set(qn("w:val"), "single" if s in sides else "nil"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), hexc); b.append(e)
    p.append(b)

def build(key, xlsx, narrative=None, out=None):
    P = PROFILES[key]; N = narrative or {}
    title, ndays, hpd, rows = read(xlsx)
    month = title.split(": ")[-1]
    hrs = sum(r["hrs"] for r in rows); bud = sum(r["bud"] for r in rows); prior = sum(r["prior"] for r in rows)
    pct = (prior + hrs) / bud if bud else None
    phase = N.get("phase") or next((r["phase"] for r in rows if r["phase"]), "the current phase")
    RGB = lambda h: RGBColor.from_string(h)

    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.75); sec.top_margin = Inches(0.6); sec.bottom_margin = Inches(0.6)
    st = doc.styles["Normal"]; st.font.name = P["sans"]; st.font.size = Pt(10); st.font.color.rgb = RGB(P["ink"])
    st.element.rPr.rFonts.set(qn("w:eastAsia"), P["sans"])

    def para(text="", font=None, size=10, bold=False, color=None, after=4, before=0, italic=False):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
        if text:
            r = p.add_run(text); r.font.name = font or P["sans"]; r.font.size = Pt(size); r.bold = bold; r.italic = italic
            r.font.color.rgb = RGB(color or P["ink"])
        return p
    def rule(p, color):
        pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "8"), ("space", "4"), ("color", color)): e.set(qn(f"w:{k}"), v)
        b.append(e); pPr.append(b)
    def heading(t):
        p = para(t, P["serif"], 13, key == "sikich", P["ink"] if key == "upstart" else P["accent"], 3, 8)
    def bullets(items, ph):
        for t in (items or [ph]):
            p = para(after=1); p.paragraph_format.left_indent = Inches(0.2); p.paragraph_format.first_line_indent = Inches(-0.15)
            r = p.add_run("•  " + t); r.font.size = Pt(9.5)
            if not items: r.font.color.rgb = RGB(P["muted"]); r.italic = True

    para(P["tag"].replace("  |  ", "   /   "), P["mono"], 8, False, P["muted"], 2)
    para(f"Executive summary: {month}", P["serif"], 22 if key == "upstart" else 18, key == "sikich", P["ink"], 2)
    p = para(after=6); rule(p, P["accent"])

    t = doc.add_table(rows=2, cols=4); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    stats = [(f"{hrs:g}", "Hours billed"), (f"{hrs / hpd:.2f}", f"Days billed ({hpd:g} hr day)"),
             (f"{pct:.0%}" if pct is not None else "[enter SOW budget]", f"Through {phase}"), (str(ndays), "Working days")]
    for i, (v, l) in enumerate(stats):
        a, b = t.cell(0, i), t.cell(1, i)
        for c in (a, b): shade(c, P["band"]); c.width = Inches(1.75)
        a.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = a.paragraphs[0].add_run(v); r.font.name = P["serif"]; r.font.size = Pt(20 if "[" not in v else 11); r.bold = key == "sikich"
        r.font.color.rgb = RGB(P["emph"] if i == 2 else P["ink"])
        r = b.paragraphs[0].add_run(l); r.font.name = P["mono"]; r.font.size = Pt(8); r.font.color.rgb = RGB(P["muted"])
        b.paragraphs[0].paragraph_format.space_after = Pt(2)

    heading("Position against the SOW")
    if pct is not None:
        txt = (f"{prior + hrs:g} of {bud:g} budgeted hours are used ({pct:.0%}) in {phase}. This month added {hrs:g} hours "
               f"({hrs / hpd:.1f} days) to {prior:g} hours billed in prior months, leaving {max(bud - prior - hrs, 0):g} hours.")
        para(txt, size=10, after=2)
    else:
        para("[Enter the SOW budget hours and prior billed hours in the timesheet to show position against the SOW.]", size=10, color=P["muted"], italic=True, after=2)

    heading("Deliverables progress")
    tb = doc.add_table(rows=1, cols=5); widths = [2.9, 0.95, 0.9, 0.95, 0.8]
    for i, h in enumerate(["Deliverable", "Status", "Hours this month", "Hours to date", "% of budget"]):
        c = tb.rows[0].cells[i]; shade(c, P["head_fill"]); c.width = Inches(widths[i])
        r = c.paragraphs[0].add_run(h); r.font.size = Pt(8.5); r.bold = True; r.font.color.rgb = RGB(P["head_text"])
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
    for r_ in rows or [dict(name="[Add deliverables in the timesheet]", status="", hrs=0, prior=0, bud=0)]:
        cells = tb.add_row().cells; td = r_["prior"] + r_["hrs"]
        vals = [r_["name"], r_["status"], f"{r_['hrs']:g}", f"{td:g}", f"{td / r_['bud']:.0%}" if r_["bud"] else ""]
        for i, v in enumerate(vals):
            cells[i].width = Inches(widths[i]); borders(cells[i], P["line"]); pp = cells[i].paragraphs[0]
            pp.paragraph_format.space_after = Pt(0); rr = pp.add_run(v); rr.font.size = Pt(9)
            if i: pp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading("Risks and decisions needed"); bullets(N.get("risks"), "[Add risks, blockers or decisions needed from the client.]")
    heading("Next month"); bullets(N.get("next"), "[Add planned work and expected hours for next month.]")

    p = para(f"Prepared by Alexis Brochu. Hours are from the {month} timesheet.", P["mono"], 8, False, P["muted"], 0, 10)
    out = out or xlsx.replace("Timesheet", "Executive-Summary").replace(".xlsx", ".docx"); doc.save(out); return out

if __name__ == "__main__":
    nar = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else None
    print(build(sys.argv[1], sys.argv[2], nar))
